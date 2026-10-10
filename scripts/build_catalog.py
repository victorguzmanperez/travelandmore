from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import unicodedata

from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook

try:
    from PIL import Image, ImageOps, UnidentifiedImageError
except ImportError:  # pragma: no cover - mensaje operativo en runtime
    Image = None
    ImageOps = None
    UnidentifiedImageError = OSError


ROOT = Path(__file__).resolve().parents[1]

SOURCE = ROOT / "content" / "Travel_More_Catalogo.xlsx"

BOOKS_OUTPUT = ROOT / "assets" / "data" / "books.js"
PRODUCTS_OUTPUT = ROOT / "assets" / "data" / "products.js"

PRODUCT_IMAGE_SOURCE_PREFIX = "assets/img/products/"
PRODUCT_IMAGE_GENERATED_PREFIX = "assets/img/products/generated/"
PRODUCT_IMAGE_SIZE = (1200, 900)
PRODUCT_IMAGE_INNER_SIZE = (1100, 800)
PRODUCT_IMAGE_BACKGROUND = (255, 253, 249)
PRODUCT_IMAGE_WEBP_QUALITY = 88


AMAZON_LINK_TIMEOUT = 12.0
AMAZON_MAX_RESPONSE_BYTES = 750_000

AMAZON_INACTIVE_MARKERS = (
    "estas buscando algo",
    "la direccion web que has especificado no es una pagina activa de nuestro sitio",
    "looking for something",
    "the web address you entered is not a functioning page on our site",
)

AMAZON_BLOCKED_MARKERS = (
    "enter the characters you see below",
    "sorry, we just need to make sure you're not a robot",
    "introduce los caracteres que ves a continuacion",
)


SHEETS = {
    "books": "01_LIBROS",
    "products": "02_COMPRAS",
    "book_categories": "03_CATEGORIAS_LIBROS",
    "product_categories": "04_CATEGORIAS_COMPRAS",
}


BOOK_HEADERS = [
    "id",
    "activo",
    "titulo",
    "autor",
    "categoria",
    "estado",
    "imagen",
    "descripcion",
    "opinion",
    "rating",
    "substack_url",
    "affiliate_url",
    "destacado",
    "orden",
]


PRODUCT_HEADERS = [
    "id",
    "activo",
    "nombre",
    "categoria",
    "imagen",
    "descripcion",
    "uso",
    "lo_que_nos_gusta",
    "lo_que_no",
    "opinion",
    "rating",
    "affiliate_url",
    "destacado",
    "orden",
]


CATEGORY_HEADERS = [
    "id",
    "categoria",
    "orden",
    "activa",
    "descripcion",
]


YES = "S\u00ed"
NO = "No"

BOOK_STATES = {
    "Le\u00eddo",
    "Leyendo",
    "Pendiente",
}


ID_PATTERN = re.compile(
    r"^[a-z0-9]+(?:-[a-z0-9]+)*$"
)


class CatalogValidationError(Exception):
    pass


def cell_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_path(value):
    return cell_text(value).replace("\\", "/")


def row_is_empty(values):
    return all(
        cell_text(value) == ""
        for value in values
    )


def read_table(ws, expected_headers, errors):
    actual_headers = [
        cell_text(cell.value)
        for cell in ws[1]
    ]

    actual_headers = actual_headers[:len(expected_headers)]

    if actual_headers != expected_headers:
        errors.append(
            f"{ws.title}: cabeceras incorrectas.\n"
            f"  Esperadas: {expected_headers}\n"
            f"  Encontradas: {actual_headers}"
        )
        return []

    rows = []

    for row_number, row in enumerate(
        ws.iter_rows(
            min_row=2,
            max_col=len(expected_headers),
            values_only=True,
        ),
        start=2,
    ):
        if row_is_empty(row):
            continue

        data = dict(
            zip(
                expected_headers,
                row,
            )
        )

        rows.append(
            (
                row_number,
                data,
            )
        )

    return rows


def parse_yes_no(
    value,
    field,
    location,
    errors,
):
    text = cell_text(value)

    if text == YES:
        return True

    if text == NO:
        return False

    errors.append(
        f"{location}: {field} debe ser "
        f"'{YES}' o '{NO}'."
    )

    return False


def validate_id(
    value,
    field,
    location,
    errors,
):
    item_id = cell_text(value)

    if not item_id:
        errors.append(
            f"{location}: {field} es obligatorio."
        )
        return ""

    if not ID_PATTERN.fullmatch(item_id):
        errors.append(
            f"{location}: {field} '{item_id}' no es valido. "
            "Usa minusculas, numeros y guiones."
        )

    return item_id


def parse_order(
    value,
    location,
    errors,
):
    if value is None or cell_text(value) == "":
        return 0

    if isinstance(value, bool):
        errors.append(
            f"{location}: orden debe ser un entero >= 0."
        )
        return 0

    try:
        number = int(value)
    except (TypeError, ValueError):
        errors.append(
            f"{location}: orden debe ser un entero >= 0."
        )
        return 0

    try:
        if float(value) != number:
            raise ValueError
    except (TypeError, ValueError):
        errors.append(
            f"{location}: orden debe ser un entero >= 0."
        )
        return 0

    if number < 0:
        errors.append(
            f"{location}: orden no puede ser negativo."
        )

    return number


def parse_rating(
    value,
    location,
    errors,
):
    if value is None or cell_text(value) == "":
        return None

    if isinstance(value, bool):
        errors.append(
            f"{location}: rating debe estar entre 1 y 5."
        )
        return None

    try:
        number = int(value)
    except (TypeError, ValueError):
        errors.append(
            f"{location}: rating debe estar entre 1 y 5."
        )
        return None

    try:
        if float(value) != number:
            raise ValueError
    except (TypeError, ValueError):
        errors.append(
            f"{location}: rating debe ser un entero entre 1 y 5."
        )
        return None

    if not 1 <= number <= 5:
        errors.append(
            f"{location}: rating debe estar entre 1 y 5."
        )

    return number


def validate_https_url(
    value,
    field,
    location,
    errors,
):
    url = cell_text(value)

    if not url:
        return ""

    if any(
        char in url
        for char in ("\r", "\n", "\t")
    ):
        errors.append(
            f"{location}: {field} contiene caracteres no validos."
        )
        return url

    try:
        parsed = urlparse(url)
    except ValueError:
        errors.append(
            f"{location}: {field} no es una URL valida."
        )
        return url

    if parsed.scheme != "https":
        errors.append(
            f"{location}: {field} debe usar https://"
        )

    if not parsed.netloc:
        errors.append(
            f"{location}: {field} no tiene dominio."
        )

    if parsed.username or parsed.password:
        errors.append(
            f"{location}: {field} no puede incluir credenciales."
        )

    return url


def normalize_web_text(value):
    text = unicodedata.normalize(
        "NFKD",
        str(value),
    )

    text = "".join(
        char
        for char in text
        if not unicodedata.combining(char)
    )

    return " ".join(
        text.casefold().split()
    )


def is_amazon_url(url):
    try:
        host = (
            urlparse(url).hostname
            or ""
        ).casefold()
    except ValueError:
        return False

    return (
        host == "link.amazon"
        or host == "amzn.eu"
        or host == "amazon.es"
        or host.endswith(".amazon.es")
    )


def classify_amazon_response(
    status_code,
    final_url,
    body,
):
    normalized_body = normalize_web_text(
        body
    )

    normalized_url = normalize_web_text(
        final_url
    )

    if status_code in (404, 410):
        return (
            "broken",
            f"HTTP {status_code}",
        )

    if "cs_404_link" in normalized_url:
        return (
            "broken",
            "Amazon redirected to its inactive-product page",
        )

    for marker in AMAZON_INACTIVE_MARKERS:
        if marker in normalized_body:
            return (
                "broken",
                "Amazon says the page is no longer active",
            )

    if status_code in (401, 403, 429, 503):
        return (
            "unverified",
            f"Amazon returned HTTP {status_code}",
        )

    for marker in AMAZON_BLOCKED_MARKERS:
        if marker in normalized_body:
            return (
                "unverified",
                "Amazon displayed an anti-bot check",
            )

    if 200 <= status_code < 400:
        return (
            "valid",
            f"HTTP {status_code}",
        )

    return (
        "unverified",
        f"HTTP {status_code}",
    )


def check_amazon_url(
    url,
    timeout=AMAZON_LINK_TIMEOUT,
):
    request = Request(
        url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/131.0 Safari/537.36"
            ),
            "Accept-Language": (
                "es-ES,es;q=0.9,en;q=0.7"
            ),
        },
    )

    try:
        with urlopen(
            request,
            timeout=timeout,
        ) as response:
            status_code = response.getcode()
            final_url = response.geturl()

            payload = response.read(
                AMAZON_MAX_RESPONSE_BYTES
            )

    except HTTPError as exc:
        if exc.code in (404, 410):
            return {
                "status": "broken",
                "reason": f"HTTP {exc.code}",
                "final_url": (vars(exc).get("filename") or url),
            }

        return {
            "status": "unverified",
            "reason": (
                f"Amazon returned HTTP {exc.code}"
            ),
            "final_url": (vars(exc).get("filename") or url),
        }

    except (
        URLError,
        TimeoutError,
        OSError,
    ) as exc:
        return {
            "status": "unverified",
            "reason": f"network error: {exc}",
            "final_url": url,
        }

    body = payload.decode(
        "utf-8",
        errors="replace",
    )

    status, reason = classify_amazon_response(
        status_code,
        final_url,
        body,
    )

    return {
        "status": status,
        "reason": reason,
        "final_url": final_url,
    }


def validate_amazon_links(
    books,
    products,
    timeout=AMAZON_LINK_TIMEOUT,
):
    targets = []

    for item in books:
        url = item.get(
            "affiliate_url",
            "",
        )

        if url and is_amazon_url(url):
            targets.append(
                (
                    "libro",
                    item["id"],
                    item["titulo"],
                    url,
                )
            )

    for item in products:
        url = item.get(
            "affiliate_url",
            "",
        )

        if url and is_amazon_url(url):
            targets.append(
                (
                    "producto",
                    item["id"],
                    item["nombre"],
                    url,
                )
            )

    broken = []
    warnings = []
    valid_count = 0

    print(
        "AMAZON LINKS: "
        f"checking {len(targets)} informed link(s)..."
    )

    for (
        kind,
        item_id,
        label,
        url,
    ) in targets:

        result = check_amazon_url(
            url,
            timeout=timeout,
        )

        status = result["status"]
        reason = result["reason"]

        if status == "valid":
            valid_count += 1

            print(
                f"PASS LINK: {kind} '{label}'"
            )

            continue

        if status == "broken":
            message = (
                f"{kind} '{label}' ({item_id}): "
                f"inactive affiliate_url: {reason}. "
                f"URL: {url}"
            )

            broken.append(
                message
            )

            print(
                f"BROKEN LINK: "
                f"{kind} '{label}' -> {reason}"
            )

            continue

        message = (
            f"{kind} '{label}' ({item_id}): "
            f"could not verify: {reason}. "
            f"URL: {url}"
        )

        warnings.append(
            message
        )

        print(
            f"WARN LINK: "
            f"{kind} '{label}' -> {reason}"
        )

    print(
        "AMAZON LINKS: "
        f"{valid_count} valid, "
        f"{len(broken)} broken, "
        f"{len(warnings)} unverified."
    )

    if warnings:
        print(
            "WARN: unverified links are not classified "
            "as broken. Review them manually before "
            "publishing if the warning persists."
        )

    return (
        broken,
        warnings,
    )


def validate_image(
    value,
    expected_prefix,
    location,
    errors,
):
    image = normalize_path(value)

    if not image:
        return ""

    path = Path(image)

    if path.is_absolute():
        errors.append(
            f"{location}: imagen debe ser una ruta relativa."
        )
        return image

    if ".." in path.parts:
        errors.append(
            f"{location}: imagen no puede contener '..'."
        )

    if not image.startswith(expected_prefix):
        errors.append(
            f"{location}: imagen debe comenzar por "
            f"'{expected_prefix}'."
        )

    full_path = ROOT / image

    if not full_path.is_file():
        errors.append(
            f"{location}: no existe la imagen '{image}'."
        )

    return image


def product_generated_relative_path(item_id):
    return (
        f"{PRODUCT_IMAGE_GENERATED_PREFIX}"
        f"{item_id}.webp"
    )


def classify_orientation(width, height):
    if width <= 0 or height <= 0:
        return "invalid"

    ratio = width / height

    if 0.95 <= ratio <= 1.05:
        return "square"

    if ratio > 1:
        return "horizontal"

    return "vertical"


def normalize_product_image(source_path):
    if Image is None or ImageOps is None:
        raise CatalogValidationError(
            "Falta Pillow para procesar imagenes de producto. "
            "Instala dependencias con: "
            "python -m pip install -r requirements.txt"
        )

    try:
        with Image.open(source_path) as opened:
            source = ImageOps.exif_transpose(opened)
            source.load()

            width, height = source.size

            if width <= 0 or height <= 0:
                raise ValueError(
                    "dimensiones de imagen no validas"
                )

            orientation = classify_orientation(
                width,
                height,
            )

            rgba = source.convert("RGBA")

    except (
        OSError,
        ValueError,
        UnidentifiedImageError,
    ) as exc:
        raise ValueError(
            f"no se puede abrir/procesar la imagen: {exc}"
        ) from exc

    inner_width, inner_height = (
        PRODUCT_IMAGE_INNER_SIZE
    )

    scale = min(
        inner_width / width,
        inner_height / height,
    )

    target_width = max(
        1,
        round(width * scale),
    )

    target_height = max(
        1,
        round(height * scale),
    )

    resized = rgba.resize(
        (target_width, target_height),
        Image.Resampling.LANCZOS,
    )

    canvas = Image.new(
        "RGBA",
        PRODUCT_IMAGE_SIZE,
        (*PRODUCT_IMAGE_BACKGROUND, 255),
    )

    left = (
        PRODUCT_IMAGE_SIZE[0]
        - target_width
    ) // 2

    top = (
        PRODUCT_IMAGE_SIZE[1]
        - target_height
    ) // 2

    canvas.alpha_composite(
        resized,
        (left, top),
    )

    output = io.BytesIO()

    canvas.convert("RGB").save(
        output,
        format="WEBP",
        quality=PRODUCT_IMAGE_WEBP_QUALITY,
        method=6,
    )

    return (
        output.getvalue(),
        orientation,
    )


def prepare_product_images(
    products,
    errors,
):
    outputs = {}

    stats = {
        "vertical": 0,
        "horizontal": 0,
        "square": 0,
    }

    for product in products:
        source_relative = product["imagen"]

        if not source_relative:
            continue

        location = (
            f"producto '{product['id']}'"
        )

        if source_relative.startswith(
            PRODUCT_IMAGE_GENERATED_PREFIX
        ):
            errors.append(
                f"{location}: imagen debe apuntar al original "
                f"en '{PRODUCT_IMAGE_SOURCE_PREFIX}', no a "
                "la carpeta generated."
            )
            continue

        source_path = (
            ROOT
            / source_relative
        )

        try:
            payload, orientation = (
                normalize_product_image(
                    source_path
                )
            )
        except (
            CatalogValidationError,
            ValueError,
        ) as exc:
            errors.append(
                f"{location}: {exc}"
            )
            continue

        if orientation in stats:
            stats[orientation] += 1

        generated_relative = (
            product_generated_relative_path(
                product["id"]
            )
        )

        generated_path = (
            ROOT
            / generated_relative
        )

        outputs[generated_path] = payload
        product["imagen"] = generated_relative

    return outputs, stats


def require_text(
    value,
    field,
    location,
    errors,
):
    text = cell_text(value)

    if not text:
        errors.append(
            f"{location}: {field} es obligatorio."
        )

    return text


def load_categories(
    rows,
    sheet_name,
    errors,
):
    categories = []
    by_name = {}

    seen_ids = set()
    seen_names = set()

    for row_number, data in rows:
        location = f"{sheet_name}!fila {row_number}"

        category_id = validate_id(
            data["id"],
            "id",
            location,
            errors,
        )

        name = require_text(
            data["categoria"],
            "categoria",
            location,
            errors,
        )

        active = parse_yes_no(
            data["activa"],
            "activa",
            location,
            errors,
        )

        order = parse_order(
            data["orden"],
            location,
            errors,
        )

        description = cell_text(
            data["descripcion"]
        )

        if category_id:
            if category_id in seen_ids:
                errors.append(
                    f"{location}: id de categoria duplicado "
                    f"'{category_id}'."
                )
            seen_ids.add(category_id)

        if name:
            if name in seen_names:
                errors.append(
                    f"{location}: categoria duplicada '{name}'."
                )
            seen_names.add(name)

        category = {
            "id": category_id,
            "categoria": name,
            "orden": order,
            "activa": active,
            "descripcion": description,
        }

        categories.append(category)

        if name:
            by_name[name] = category

    categories.sort(
        key=lambda item: (
            item["orden"],
            item["categoria"].casefold(),
            item["id"],
        )
    )

    return categories, by_name


def validate_category_reference(
    value,
    category_map,
    location,
    errors,
):
    name = cell_text(value)

    if not name:
        errors.append(
            f"{location}: categoria es obligatoria."
        )
        return "", ""

    category = category_map.get(name)

    if category is None:
        errors.append(
            f"{location}: categoria desconocida '{name}'."
        )
        return name, ""

    if not category["activa"]:
        errors.append(
            f"{location}: categoria '{name}' esta inactiva."
        )

    return name, category["id"]


def load_books(
    rows,
    category_map,
    errors,
):
    books = []
    seen_ids = set()

    for row_number, data in rows:
        location = (
            f"{SHEETS['books']}!fila {row_number}"
        )

        item_id = validate_id(
            data["id"],
            "id",
            location,
            errors,
        )

        if item_id:
            if item_id in seen_ids:
                errors.append(
                    f"{location}: id duplicado '{item_id}'."
                )
            seen_ids.add(item_id)

        active = parse_yes_no(
            data["activo"],
            "activo",
            location,
            errors,
        )

        featured = parse_yes_no(
            data["destacado"],
            "destacado",
            location,
            errors,
        )

        order = parse_order(
            data["orden"],
            location,
            errors,
        )

        rating = parse_rating(
            data["rating"],
            location,
            errors,
        )

        substack_url = validate_https_url(
            data["substack_url"],
            "substack_url",
            location,
            errors,
        )

        affiliate_url = validate_https_url(
            data["affiliate_url"],
            "affiliate_url",
            location,
            errors,
        )

        image = validate_image(
            data["imagen"],
            "assets/img/books/",
            location,
            errors,
        )

        title = cell_text(data["titulo"])
        author = cell_text(data["autor"])
        category_name = cell_text(data["categoria"])
        state = cell_text(data["estado"])

        category_id = ""

        if category_name:
            category = category_map.get(category_name)

            if category is None:
                errors.append(
                    f"{location}: categoria desconocida "
                    f"'{category_name}'."
                )
            else:
                category_id = category["id"]

                if active and not category["activa"]:
                    errors.append(
                        f"{location}: categoria "
                        f"'{category_name}' esta inactiva."
                    )

        if active:
            title = require_text(
                data["titulo"],
                "titulo",
                location,
                errors,
            )

            author = require_text(
                data["autor"],
                "autor",
                location,
                errors,
            )

            category_name, category_id = (
                validate_category_reference(
                    data["categoria"],
                    category_map,
                    location,
                    errors,
                )
            )

            state = require_text(
                data["estado"],
                "estado",
                location,
                errors,
            )

            if (
                state
                and state not in BOOK_STATES
            ):
                errors.append(
                    f"{location}: estado '{state}' no valido."
                )

        elif state and state not in BOOK_STATES:
            errors.append(
                f"{location}: estado '{state}' no valido."
            )

        book = {
            "id": item_id,
            "titulo": title,
            "autor": author,
            "categoria": category_name,
            "categoria_id": category_id,
            "estado": state,
            "imagen": image,
            "descripcion": cell_text(
                data["descripcion"]
            ),
            "opinion": cell_text(
                data["opinion"]
            ),
            "rating": rating,
            "substack_url": substack_url,
            "affiliate_url": affiliate_url,
            "destacado": featured,
            "orden": order,
        }

        if active:
            books.append(book)

    books.sort(
        key=lambda item: (
            item["orden"],
            item["titulo"].casefold(),
            item["id"],
        )
    )

    return books


def load_products(
    rows,
    category_map,
    errors,
):
    products = []
    seen_ids = set()

    for row_number, data in rows:
        location = (
            f"{SHEETS['products']}!fila {row_number}"
        )

        item_id = validate_id(
            data["id"],
            "id",
            location,
            errors,
        )

        if item_id:
            if item_id in seen_ids:
                errors.append(
                    f"{location}: id duplicado '{item_id}'."
                )
            seen_ids.add(item_id)

        active = parse_yes_no(
            data["activo"],
            "activo",
            location,
            errors,
        )

        featured = parse_yes_no(
            data["destacado"],
            "destacado",
            location,
            errors,
        )

        order = parse_order(
            data["orden"],
            location,
            errors,
        )

        rating = parse_rating(
            data["rating"],
            location,
            errors,
        )

        affiliate_url = validate_https_url(
            data["affiliate_url"],
            "affiliate_url",
            location,
            errors,
        )

        image = validate_image(
            data["imagen"],
            PRODUCT_IMAGE_SOURCE_PREFIX,
            location,
            errors,
        )

        name = cell_text(data["nombre"])
        category_name = cell_text(data["categoria"])
        category_id = ""

        if category_name:
            category = category_map.get(category_name)

            if category is None:
                errors.append(
                    f"{location}: categoria desconocida "
                    f"'{category_name}'."
                )
            else:
                category_id = category["id"]

                if active and not category["activa"]:
                    errors.append(
                        f"{location}: categoria "
                        f"'{category_name}' esta inactiva."
                    )

        if active:
            name = require_text(
                data["nombre"],
                "nombre",
                location,
                errors,
            )

            category_name, category_id = (
                validate_category_reference(
                    data["categoria"],
                    category_map,
                    location,
                    errors,
                )
            )

        product = {
            "id": item_id,
            "nombre": name,
            "categoria": category_name,
            "categoria_id": category_id,
            "imagen": image,
            "descripcion": cell_text(
                data["descripcion"]
            ),
            "uso": cell_text(
                data["uso"]
            ),
            "lo_que_nos_gusta": cell_text(
                data["lo_que_nos_gusta"]
            ),
            "lo_que_no": cell_text(
                data["lo_que_no"]
            ),
            "opinion": cell_text(
                data["opinion"]
            ),
            "rating": rating,
            "affiliate_url": affiliate_url,
            "destacado": featured,
            "orden": order,
        }

        if active:
            products.append(product)

    products.sort(
        key=lambda item: (
            item["orden"],
            item["nombre"].casefold(),
            item["id"],
        )
    )

    return products


def public_categories(
    categories,
):
    return [
        {
            "id": item["id"],
            "categoria": item["categoria"],
            "orden": item["orden"],
            "descripcion": item["descripcion"],
        }
        for item in categories
        if item["activa"]
    ]


def js_export(
    export_name,
    value,
):
    payload = json.dumps(
        value,
        ensure_ascii=False,
        indent=2,
    )

    return (
        f"export const {export_name} = "
        f"{payload};\n"
    )


def build_books_js(
    categories,
    books,
):
    header = (
        "// AUTO-GENERATED by scripts/build_catalog.py.\n"
        "// Source: content/Travel_More_Catalogo.xlsx\n"
        "// DO NOT EDIT MANUALLY.\n\n"
    )

    return (
        header
        + js_export(
            "bookCategories",
            public_categories(categories),
        )
        + "\n"
        + js_export(
            "books",
            books,
        )
    )


def build_products_js(
    categories,
    products,
):
    header = (
        "// AUTO-GENERATED by scripts/build_catalog.py.\n"
        "// Source: content/Travel_More_Catalogo.xlsx\n"
        "// DO NOT EDIT MANUALLY.\n\n"
    )

    return (
        header
        + js_export(
            "productCategories",
            public_categories(categories),
        )
        + "\n"
        + js_export(
            "products",
            products,
        )
    )


def transactional_write(outputs):
    old_contents = {}

    temp_paths = {}

    for path, content in outputs.items():
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        old_contents[path] = (
            path.read_bytes()
            if path.exists()
            else None
        )

        temp = path.with_suffix(
            path.suffix + ".tmp"
        )

        payload = (
            content.encode("utf-8")
            if isinstance(content, str)
            else bytes(content)
        )

        temp.write_bytes(
            payload
        )

        temp_paths[path] = temp

    replaced = []

    try:
        for path, temp in temp_paths.items():
            os.replace(
                temp,
                path,
            )

            replaced.append(path)

    except Exception:
        for path in replaced:
            old = old_contents[path]

            if old is None:
                try:
                    path.unlink()
                except FileNotFoundError:
                    pass
            else:
                path.write_bytes(old)

        raise

    finally:
        for temp in temp_paths.values():
            try:
                temp.unlink()
            except FileNotFoundError:
                pass


def validate_source_exists():
    if not SOURCE.is_file():
        raise CatalogValidationError(
            "No existe "
            "content/Travel_More_Catalogo.xlsx"
        )


def build(
    check_only=False,
    check_links=False,
    link_timeout=AMAZON_LINK_TIMEOUT,
):
    validate_source_exists()

    errors = []

    workbook = load_workbook(
        SOURCE,
        data_only=True,
        read_only=False,
    )

    required_sheets = [
        "00_CONFIG",
        SHEETS["books"],
        SHEETS["products"],
        SHEETS["book_categories"],
        SHEETS["product_categories"],
    ]

    for sheet in required_sheets:
        if sheet not in workbook.sheetnames:
            errors.append(
                f"Falta la hoja obligatoria '{sheet}'."
            )

    if errors:
        raise CatalogValidationError(
            "\n".join(errors)
        )

    book_category_rows = read_table(
        workbook[
            SHEETS["book_categories"]
        ],
        CATEGORY_HEADERS,
        errors,
    )

    product_category_rows = read_table(
        workbook[
            SHEETS["product_categories"]
        ],
        CATEGORY_HEADERS,
        errors,
    )

    book_categories, book_category_map = (
        load_categories(
            book_category_rows,
            SHEETS["book_categories"],
            errors,
        )
    )

    product_categories, product_category_map = (
        load_categories(
            product_category_rows,
            SHEETS["product_categories"],
            errors,
        )
    )

    book_rows = read_table(
        workbook[SHEETS["books"]],
        BOOK_HEADERS,
        errors,
    )

    product_rows = read_table(
        workbook[SHEETS["products"]],
        PRODUCT_HEADERS,
        errors,
    )

    books = load_books(
        book_rows,
        book_category_map,
        errors,
    )

    products = load_products(
        product_rows,
        product_category_map,
        errors,
    )

    product_image_outputs, product_image_stats = (
        prepare_product_images(
            products,
            errors,
        )
    )

    if errors:
        message = [
            "CATALOGO INVALIDO",
            "",
        ]

        message.extend(
            f"- {error}"
            for error in errors
        )

        raise CatalogValidationError(
            "\n".join(message)
        )

    if check_links:
        broken_links, _warnings = validate_amazon_links(
            books,
            products,
            timeout=link_timeout,
        )

        if broken_links:
            message = [
                "ENLACES AMAZON INVALIDOS",
                "",
            ]

            message.extend(
                f"- {error}"
                for error in broken_links
            )

            raise CatalogValidationError(
                "\n".join(message)
            )

    books_js = build_books_js(
        book_categories,
        books,
    )

    products_js = build_products_js(
        product_categories,
        products,
    )

    print(
        f"PASS: {len(books)} libro(s) activo(s)"
    )

    print(
        f"PASS: {len(products)} producto(s) activo(s)"
    )

    print(
        "PASS: "
        f"{len(public_categories(book_categories))} "
        "categoria(s) de libros activa(s)"
    )

    print(
        "PASS: "
        f"{len(public_categories(product_categories))} "
        "categoria(s) de compras activa(s)"
    )

    image_count = sum(
        product_image_stats.values()
    )

    print(
        "PASS: "
        f"{image_count} imagen(es) de producto valida(s) "
        f"({product_image_stats['vertical']} vertical(es), "
        f"{product_image_stats['horizontal']} horizontal(es), "
        f"{product_image_stats['square']} cuadrada(s))"
    )

    if check_only:
        print(
            "CHECK ONLY: no se han modificado "
            "books.js, products.js ni imagenes generated."
        )

        return

    outputs = {
        BOOKS_OUTPUT: books_js,
        PRODUCTS_OUTPUT: products_js,
    }

    outputs.update(
        product_image_outputs
    )

    transactional_write(
        outputs
    )

    print(
        "PASS: assets/data/books.js generado"
    )

    print(
        "PASS: "
        f"{len(product_image_outputs)} imagen(es) de producto "
        "generada(s) en assets/img/products/generated/"
    )

    print(
        "PASS: assets/data/products.js generado"
    )


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Valida Travel_More_Catalogo.xlsx "
            "y genera los catalogos JS."
        )
    )

    parser.add_argument(
        "--check",
        action="store_true",
        help=(
            "Valida el Excel sin modificar "
            "books.js ni products.js."
        ),
    )

    parser.add_argument(
        "--check-links",
        action="store_true",
        help=(
            "Checks Amazon affiliate_url links online "
            "for active books and products. "
            "Does not generate files."
        ),
    )

    parser.add_argument(
        "--link-timeout",
        type=float,
        default=AMAZON_LINK_TIMEOUT,
        help=(
            "Timeout per Amazon link in seconds "
            f"(default: {AMAZON_LINK_TIMEOUT:g})."
        ),
    )

    return parser.parse_args()


def main():
    args = parse_args()

    try:
        if args.link_timeout <= 0:
            raise CatalogValidationError(
                "--link-timeout must be greater than 0."
            )

        build(
            check_only=(
                args.check
                or args.check_links
            ),
            check_links=args.check_links,
            link_timeout=args.link_timeout,
        )

    except CatalogValidationError as exc:
        print(
            "",
            file=sys.stderr,
        )

        print(
            exc,
            file=sys.stderr,
        )

        print(
            "",
            file=sys.stderr,
        )

        print(
            "FAIL-CLOSED: no se han actualizado "
            "los catalogos generados.",
            file=sys.stderr,
        )

        return 1

    except Exception as exc:
        print(
            f"ERROR INESPERADO: {exc}",
            file=sys.stderr,
        )

        print(
            "FAIL-CLOSED: no se han actualizado "
            "los catalogos generados.",
            file=sys.stderr,
        )

        return 2

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
