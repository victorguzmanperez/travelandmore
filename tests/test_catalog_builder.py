from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import re
import shutil
import tempfile
import unittest
import warnings

from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]

SOURCE = (
    ROOT
    / "content"
    / "Travel_More_Catalogo.xlsx"
)

BUILDER_PATH = (
    ROOT
    / "scripts"
    / "build_catalog.py"
)


warnings.filterwarnings(
    "ignore",
    message=(
        "Data Validation extension "
        "is not supported and will be removed"
    ),
)


spec = importlib.util.spec_from_file_location(
    "travel_more_catalog_builder",
    BUILDER_PATH,
)

builder = importlib.util.module_from_spec(spec)

assert spec.loader is not None

spec.loader.exec_module(builder)


def sha256(path):
    return hashlib.sha256(
        path.read_bytes()
    ).hexdigest()


def first_data_row(ws):
    for row in range(
        2,
        ws.max_row + 1,
    ):
        value = ws.cell(
            row=row,
            column=1,
        ).value

        if value not in (
            None,
            "",
        ):
            return row

    raise AssertionError(
        f"No hay datos en {ws.title}"
    )


def extract_export(
    text,
    export_name,
):
    pattern = (
        rf"export const "
        rf"{re.escape(export_name)}"
        rf" = (\[.*?\]);"
    )

    match = re.search(
        pattern,
        text,
        re.S,
    )

    if match is None:
        raise AssertionError(
            f"No se encuentra export {export_name}"
        )

    return json.loads(
        match.group(1)
    )


class CatalogBuilderQA(unittest.TestCase):

    def setUp(self):
        self.temp = (
            tempfile.TemporaryDirectory()
        )

        self.temp_root = Path(
            self.temp.name
        )

        self.source = (
            self.temp_root
            / "Travel_More_Catalogo.xlsx"
        )

        shutil.copy2(
            SOURCE,
            self.source,
        )

        shutil.copytree(
            ROOT / "assets" / "img",
            self.temp_root / "assets" / "img",
        )

        self.books_output = (
            self.temp_root
            / "books.js"
        )

        self.products_output = (
            self.temp_root
            / "products.js"
        )

        self.original_root = (
            builder.ROOT
        )

        self.original_source = (
            builder.SOURCE
        )

        self.original_books_output = (
            builder.BOOKS_OUTPUT
        )

        self.original_products_output = (
            builder.PRODUCTS_OUTPUT
        )

        builder.ROOT = self.temp_root
        builder.SOURCE = self.source
        builder.BOOKS_OUTPUT = (
            self.books_output
        )
        builder.PRODUCTS_OUTPUT = (
            self.products_output
        )

    def tearDown(self):
        builder.ROOT = (
            self.original_root
        )

        builder.SOURCE = (
            self.original_source
        )

        builder.BOOKS_OUTPUT = (
            self.original_books_output
        )

        builder.PRODUCTS_OUTPUT = (
            self.original_products_output
        )

        self.temp.cleanup()

    def mutate(self, callback):
        workbook = load_workbook(
            self.source
        )

        callback(workbook)

        workbook.save(
            self.source
        )

    def build(self):
        output = io.StringIO()

        with redirect_stdout(output):
            builder.build(
                check_only=False
            )

        return output.getvalue()

    def check(self):
        output = io.StringIO()

        with redirect_stdout(output):
            builder.build(
                check_only=True
            )

        return output.getvalue()

    def make_product_image(
        self,
        filename,
        size,
        color=(20, 30, 40),
    ):
        path = (
            self.temp_root
            / "assets"
            / "img"
            / "products"
            / filename
        )

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        Image.new(
            "RGB",
            size,
            color,
        ).save(path)

        return path

    def expect_invalid(
        self,
        mutator,
        expected_text,
    ):
        sentinel_books = (
            "// BOOKS_SENTINEL\n"
        )

        sentinel_products = (
            "// PRODUCTS_SENTINEL\n"
        )

        self.books_output.write_text(
            sentinel_books,
            encoding="utf-8",
        )

        self.products_output.write_text(
            sentinel_products,
            encoding="utf-8",
        )

        self.mutate(mutator)

        with self.assertRaises(
            builder.CatalogValidationError
        ) as context:
            self.build()

        self.assertIn(
            expected_text,
            str(context.exception),
        )

        self.assertEqual(
            self.books_output.read_text(
                encoding="utf-8"
            ),
            sentinel_books,
        )

        self.assertEqual(
            self.products_output.read_text(
                encoding="utf-8"
            ),
            sentinel_products,
        )

    def test_01_valid_catalog_builds(self):
        self.build()

        self.assertTrue(
            self.books_output.is_file()
        )

        self.assertTrue(
            self.products_output.is_file()
        )

        text = self.books_output.read_text(
            encoding="utf-8"
        )

        self.assertIn(
            "Padre rico, padre pobre",
            text,
        )

    def test_02_two_builds_are_identical(self):
        self.build()

        first_books = sha256(
            self.books_output
        )

        first_products = sha256(
            self.products_output
        )

        self.build()

        self.assertEqual(
            first_books,
            sha256(
                self.books_output
            ),
        )

        self.assertEqual(
            first_products,
            sha256(
                self.products_output
            ),
        )

    def test_03_duplicate_id_fails_closed(self):

        def mutate(workbook):
            ws = workbook[
                "01_LIBROS"
            ]

            source_row = first_data_row(
                ws
            )

            target_row = (
                ws.max_row + 1
            )

            for column in range(
                1,
                15,
            ):
                ws.cell(
                    row=target_row,
                    column=column,
                ).value = ws.cell(
                    row=source_row,
                    column=column,
                ).value

        self.expect_invalid(
            mutate,
            "id duplicado",
        )

    def test_04_unknown_category_fails_closed(self):

        def mutate(workbook):
            ws = workbook[
                "01_LIBROS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=5,
            ).value = (
                "Categoria QA inexistente"
            )

        self.expect_invalid(
            mutate,
            "categoria desconocida",
        )

    def test_05_invalid_rating_fails_closed(self):

        def mutate(workbook):
            ws = workbook[
                "01_LIBROS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=10,
            ).value = 8

        self.expect_invalid(
            mutate,
            "rating debe estar entre 1 y 5",
        )

    def test_06_http_url_fails_closed(self):

        def mutate(workbook):
            ws = workbook[
                "01_LIBROS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=12,
            ).value = (
                "http://example.com/qa"
            )

        self.expect_invalid(
            mutate,
            "debe usar https://",
        )

    def test_07_missing_image_fails_closed(self):

        def mutate(workbook):
            ws = workbook[
                "01_LIBROS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=7,
            ).value = (
                "assets/img/books/"
                "__qa_no_existe__.jpg"
            )

        self.expect_invalid(
            mutate,
            "no existe la imagen",
        )

    def test_08_inactive_book_is_not_published(self):

        workbook = load_workbook(
            self.source
        )

        ws = workbook[
            "01_LIBROS"
        ]

        row = first_data_row(
            ws
        )

        item_id = str(
            ws.cell(
                row=row,
                column=1,
            ).value
        )

        ws.cell(
            row=row,
            column=2,
        ).value = "No"

        workbook.save(
            self.source
        )

        self.build()

        output = (
            self.books_output.read_text(
                encoding="utf-8"
            )
        )

        self.assertNotIn(
            f'"id": "{item_id}"',
            output,
        )

    def test_09_products_pipeline_accepts_valid_item(self):

        def mutate(workbook):
            ws = workbook[
                "02_COMPRAS"
            ]

            row = 2

            values = [
                "qa-producto",
                "S\u00ed",
                "Producto QA",
                "Tecnolog\u00eda",
                "",
                "Producto temporal de QA.",
                "Uso temporal de QA.",
                "Punto positivo QA.",
                "Punto negativo QA.",
                "Opinion QA.",
                4,
                "https://example.com/qa",
                "No",
                9999,
            ]

            for index, value in enumerate(
                values,
                start=1,
            ):
                ws.cell(
                    row=row,
                    column=index,
                ).value = value

        self.mutate(
            mutate
        )

        self.build()

        output = (
            self.products_output.read_text(
                encoding="utf-8"
            )
        )

        self.assertIn(
            '"id": "qa-producto"',
            output,
        )

    def test_10_unused_categories_are_not_rendered(self):
        self.build()

        generated = (
            self.books_output.read_text(
                encoding="utf-8"
            )
        )

        categories = extract_export(
            generated,
            "bookCategories",
        )

        books = extract_export(
            generated,
            "books",
        )

        used_ids = {
            book["categoria_id"]
            for book in books
        }

        unused = [
            category
            for category in categories
            if category["id"]
            not in used_ids
        ]

        self.assertGreater(
            len(unused),
            0,
        )

        renderer = (
            ROOT
            / "assets"
            / "js"
            / "catalog.js"
        ).read_text(
            encoding="utf-8"
        )

        self.assertIn(
            "getUsedCategories",
            renderer,
        )

        self.assertIn(
            "usedCategoryIds.has(category.id)",
            renderer,
        )


    def test_11_product_images_generate_4_3_webp(self):
        self.build()

        generated_text = self.products_output.read_text(
            encoding="utf-8"
        )

        products = extract_export(
            generated_text,
            "products",
        )

        with_images = [
            product
            for product in products
            if product["imagen"]
        ]

        self.assertGreaterEqual(
            len(with_images),
            1,
        )

        for product in with_images:
            self.assertTrue(
                product["imagen"].startswith(
                    "assets/img/products/generated/"
                )
            )

            generated = (
                self.temp_root
                / product["imagen"]
            )

            self.assertTrue(
                generated.is_file()
            )

            with Image.open(generated) as image:
                self.assertEqual(
                    image.size,
                    (1200, 900),
                )

                self.assertEqual(
                    image.format,
                    "WEBP",
                )

    def test_12_normalizer_preserves_aspect_ratio_without_crop(self):
        background = Image.new(
            "RGB",
            builder.PRODUCT_IMAGE_SIZE,
            builder.PRODUCT_IMAGE_BACKGROUND,
        )

        cases = [
            ("vertical.jpg", (300, 600), 0.5),
            ("horizontal.jpg", (600, 300), 2.0),
            ("square.jpg", (400, 400), 1.0),
        ]

        for filename, size, expected_ratio in cases:
            source = self.make_product_image(
                filename,
                size,
                color=(10, 10, 10),
            )

            payload, _ = builder.normalize_product_image(
                source
            )

            generated = Image.open(
                io.BytesIO(payload)
            ).convert("RGB")

            grayscale = generated.convert("L")

            mask = grayscale.point(
                lambda value: (
                    255
                    if value < 128
                    else 0
                )
            )

            bbox = mask.getbbox()

            self.assertIsNotNone(bbox)

            left, top, right, bottom = bbox
            content_width = right - left
            content_height = bottom - top

            ratio = (
                content_width
                / content_height
            )

            self.assertAlmostEqual(
                ratio,
                expected_ratio,
                delta=0.03,
            )

    def test_13_corrupt_product_image_fails_closed(self):
        corrupt = (
            self.temp_root
            / "assets"
            / "img"
            / "products"
            / "qa-corrupt.jpg"
        )

        corrupt.write_bytes(
            b"not-an-image"
        )

        def mutate(workbook):
            ws = workbook[
                "02_COMPRAS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=2,
            ).value = "S\u00ed"

            ws.cell(
                row=row,
                column=5,
            ).value = (
                "assets/img/products/"
                "qa-corrupt.jpg"
            )

        self.expect_invalid(
            mutate,
            "no se puede abrir/procesar la imagen",
        )

    def test_14_missing_product_image_fails_closed(self):
        def mutate(workbook):
            ws = workbook[
                "02_COMPRAS"
            ]

            row = first_data_row(
                ws
            )

            ws.cell(
                row=row,
                column=2,
            ).value = "S\u00ed"

            ws.cell(
                row=row,
                column=5,
            ).value = (
                "assets/img/products/"
                "__qa_missing__.jpg"
            )

        self.expect_invalid(
            mutate,
            "no existe la imagen",
        )

    def test_15_check_only_does_not_write_generated_images(self):
        generated_dir = (
            self.temp_root
            / "assets"
            / "img"
            / "products"
            / "generated"
        )

        shutil.rmtree(
            generated_dir,
            ignore_errors=True,
        )

        output = self.check()

        self.assertIn(
            "CHECK ONLY",
            output,
        )

        self.assertFalse(
            generated_dir.exists()
        )

    def test_16_generated_product_images_are_reproducible(self):
        self.build()

        generated_dir = (
            self.temp_root
            / "assets"
            / "img"
            / "products"
            / "generated"
        )

        first = {
            path.name: sha256(path)
            for path in generated_dir.glob("*.webp")
        }

        self.assertGreaterEqual(
            len(first),
            1,
        )

        self.build()

        second = {
            path.name: sha256(path)
            for path in generated_dir.glob("*.webp")
        }

        self.assertEqual(
            first,
            second,
        )


if __name__ == "__main__":
    unittest.main(
        verbosity=2
    )



class FakeAmazonResponse:
    def __init__(
        self,
        status_code=200,
        final_url="https://www.amazon.es/dp/B000TEST",
        body="<html><body>Product available</body></html>",
    ):
        self._status_code = status_code
        self._final_url = final_url
        self._body = body.encode("utf-8")

    def __enter__(self):
        return self

    def __exit__(
        self,
        exc_type,
        exc,
        tb,
    ):
        return False

    def getcode(self):
        return self._status_code

    def geturl(self):
        return self._final_url

    def read(
        self,
        _limit=-1,
    ):
        return self._body


class AmazonLinkQA(unittest.TestCase):

    def setUp(self):
        self.original_urlopen = (
            builder.urlopen
        )

    def tearDown(self):
        builder.urlopen = (
            self.original_urlopen
        )

    def test_17_amazon_valid_200_is_valid(self):
        builder.urlopen = (
            lambda request, timeout:
            FakeAmazonResponse()
        )

        result = builder.check_amazon_url(
            "https://www.amazon.es/dp/B000TEST"
        )

        self.assertEqual(
            result["status"],
            "valid",
        )

    def test_18_amazon_http_404_is_broken(self):
        def fake_urlopen(
            request,
            timeout,
        ):
            raise builder.HTTPError(
                request.full_url,
                404,
                "Not Found",
                hdrs=None,
                fp=None,
            )

        builder.urlopen = fake_urlopen

        result = builder.check_amazon_url(
            "https://www.amazon.es/dp/B000DEAD"
        )

        self.assertEqual(
            result["status"],
            "broken",
        )

    def test_19_amazon_inactive_200_is_broken(self):
        body = """
        <html>
        <body>
        Looking for something?
        The Web address you entered is not a
        functioning page on our site.
        </body>
        </html>
        """

        builder.urlopen = (
            lambda request, timeout:
            FakeAmazonResponse(
                status_code=200,
                final_url=(
                    "https://www.amazon.es/dp/B000DEAD"
                ),
                body=body,
            )
        )

        result = builder.check_amazon_url(
            "https://www.amazon.es/dp/B000DEAD"
        )

        self.assertEqual(
            result["status"],
            "broken",
        )

    def test_20_amazon_redirect_to_product_is_valid(self):
        builder.urlopen = (
            lambda request, timeout:
            FakeAmazonResponse(
                status_code=200,
                final_url=(
                    "https://www.amazon.es/dp/B000LIVE"
                ),
                body="<html><body>Product</body></html>",
            )
        )

        result = builder.check_amazon_url(
            "https://link.amazon/B000LIVE"
        )

        self.assertEqual(
            result["status"],
            "valid",
        )

        self.assertEqual(
            result["final_url"],
            "https://www.amazon.es/dp/B000LIVE",
        )

    def test_21_amazon_network_error_is_unverified(self):
        def fake_urlopen(
            request,
            timeout,
        ):
            raise builder.URLError(
                "offline"
            )

        builder.urlopen = fake_urlopen

        result = builder.check_amazon_url(
            "https://amzn.eu/d/B000TEST"
        )

        self.assertEqual(
            result["status"],
            "unverified",
        )
